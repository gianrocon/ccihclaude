import uuid
from datetime import datetime, timedelta, date
from pathlib import Path

import openpyxl
import xlrd
from django.db import transaction, IntegrityError

from core.models import Cultura, Antibiograma, ControleAtbRaw, Internamento, Importacao
from core.services.patient_service import upsert_paciente


def _excel_serial_to_date(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    try:
        n = float(value)
        d = datetime(1899, 12, 30) + timedelta(days=n)
        return d.strftime("%Y-%m-%d")
    except Exception:
        return str(value)


def _os_str(value) -> str:
    try:
        return str(int(float(str(value).strip())))
    except Exception:
        return str(value).strip()


def _parse_date(s):
    try:
        return date.fromisoformat(s) if s else None
    except Exception:
        return None


def _detect_type(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(max_row=3, values_only=True):
            for cell in row:
                if cell and "microbiologia" in str(cell).lower():
                    wb.close()
                    return "microbiologia"
                if cell and "antimicrobiano" in str(cell).lower():
                    wb.close()
                    return "controle_atb"
                if cell and ("permanência" in str(cell).lower() or "internamento" in str(cell).lower()):
                    wb.close()
                    return "passagens"
        wb.close()
    elif ext == ".xls":
        wb = xlrd.open_workbook(str(path), encoding_override="cp1252")
        ws = wb.sheet_by_index(0)
        for i in range(min(4, ws.nrows)):
            for j in range(ws.ncols):
                val = str(ws.cell_value(i, j)).lower()
                if "microbiologia" in val:
                    return "microbiologia"
                if "antimicrobiano" in val:
                    return "controle_atb"
                if "permanência" in val or "internamento" in val:
                    return "passagens"
    return "desconhecido"


def importar_microbiologia(path: Path, hospital, usuario=None) -> int:
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx = None
    for i, row in enumerate(rows):
        if any("prontu" in str(c).lower() for c in row if c):
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Cabeçalho não encontrado no arquivo de microbiologia.")

    headers = [str(c).strip() if c else "" for c in rows[header_idx]]

    def colidx(*keywords):
        for kw in keywords:
            kw_l = kw.lower()
            for i, h in enumerate(headers):
                if kw_l in h.lower():
                    return i
        return None

    idx_os     = colidx("OS", " os") or 0
    idx_pront  = colidx("prontu") or 1
    idx_nome   = colidx("nome pac", "nome") or 2
    idx_unid   = colidx("unidade") or 3
    idx_coleta = colidx("coleta") or 4
    idx_assin  = colidx("assinatura") or 5
    idx_proc   = colidx("procedimento") or 6
    idx_micro  = colidx("microrganismo") or 7
    idx_resist = colidx("resist") or 8
    idx_sensiv = colidx("sens") or 9
    idx_interm = colidx("interm") or 10
    idx_obs    = colidx("obs") or 11
    idx_trat   = colidx("trat") or 12

    if idx_sensiv == idx_resist:
        idx_sensiv = 9
    if idx_interm == idx_resist or idx_interm == idx_sensiv:
        idx_interm = 10

    def getval(row, idx):
        if idx is None or idx >= len(row):
            return ""
        v = row[idx]
        return "" if v is None else str(v).strip()

    inserted = 0

    with transaction.atomic():
        for row_vals in rows[header_idx + 1:]:
            if not any(v for v in row_vals if v):
                continue

            os_num        = _os_str(getval(row_vals, idx_os))
            prontuario    = getval(row_vals, idx_pront)
            nome          = getval(row_vals, idx_nome)
            unidade       = getval(row_vals, idx_unid)
            dt_coleta_str = _excel_serial_to_date(row_vals[idx_coleta]) if row_vals[idx_coleta] else ""
            dt_assin_str  = _excel_serial_to_date(row_vals[idx_assin]) if row_vals[idx_assin] else ""
            procedimento  = getval(row_vals, idx_proc)
            microrganismo = getval(row_vals, idx_micro)
            resist_raw    = getval(row_vals, idx_resist)
            sensiv_raw    = getval(row_vals, idx_sensiv)
            interm_raw    = getval(row_vals, idx_interm)
            obs           = getval(row_vals, idx_obs)
            trat          = getval(row_vals, idx_trat)

            if not os_num or not prontuario:
                continue

            pac = upsert_paciente(prontuario, nome, hospital)

            try:
                with transaction.atomic():
                    cultura = Cultura.objects.create(
                        paciente=pac,
                        os=os_num,
                        unidade=unidade,
                        dt_coleta=_parse_date(dt_coleta_str),
                        dt_assinatura=_parse_date(dt_assin_str),
                        procedimento=procedimento,
                        microrganismo=microrganismo,
                        obs=obs,
                        trat_respir=trat,
                    )
                    inserted += 1

                    def insert_atbs(lista_raw, sens):
                        if not lista_raw:
                            return
                        for atb in str(lista_raw).split(","):
                            atb = atb.strip()
                            if atb:
                                Antibiograma.objects.create(
                                    cultura=cultura,
                                    antibiotico=atb,
                                    sensibilidade=sens,
                                )

                    insert_atbs(resist_raw, "R")
                    insert_atbs(sensiv_raw, "S")
                    insert_atbs(interm_raw, "I")

            except IntegrityError:
                pass

    Importacao.objects.create(
        hospital=hospital, usuario=usuario,
        arquivo=path.name, tipo="microbiologia", registros=inserted,
    )
    return inserted


_XLS_COL_INICIO    = 0
_XLS_COL_DIAS_SOL  = 1
_XLS_COL_DIAS_CCIH = 3
_XLS_COL_MED       = 6
_XLS_COL_DIAS_USO  = 13
_XLS_COL_MEDICO    = 16
_XLS_COL_ACOM      = 14


def importar_controle_atb(path: Path, hospital, usuario=None) -> int:
    path = Path(path)
    ext = path.suffix.lower()

    if ext == ".xls":
        wb = xlrd.open_workbook(str(path), encoding_override="cp1252")
        ws = wb.sheet_by_index(0)
    else:
        raise ValueError(f"Use .xls para arquivo de controle ATB (recebido: {ext})")

    inserted = 0
    current_pront = None
    current_nome = ""
    current_acom = ""

    with transaction.atomic():
        for i in range(ws.nrows):
            ct0 = ws.cell_type(i, 0)
            v0  = ws.cell_value(i, 0)
            v0s = str(v0).strip()

            if "impresso em" in v0s.lower():
                break

            if ct0 == xlrd.XL_CELL_TEXT and "prontu" in v0s.lower():
                continue

            if ct0 == xlrd.XL_CELL_TEXT and ("início" in v0s.lower() or "inicio" in v0s.lower()):
                continue

            if ct0 == xlrd.XL_CELL_TEXT and ws.cell_type(i, 2) == xlrd.XL_CELL_TEXT:
                v2 = str(ws.cell_value(i, 2)).strip()
                if v2 and v2 not in ("Paciente",):
                    current_pront = v0s
                    current_nome  = v2
                    current_acom  = str(ws.cell_value(i, _XLS_COL_ACOM)).strip() if ws.ncols > _XLS_COL_ACOM else ""
                    continue

            if ct0 == xlrd.XL_CELL_DATE and current_pront:
                dt_inicio_str = _excel_serial_to_date(v0)

                def nint(col):
                    try:
                        return int(float(ws.cell_value(i, col))) if ws.ncols > col else 0
                    except Exception:
                        return 0

                def nstr(col):
                    return str(ws.cell_value(i, col)).strip() if ws.ncols > col else ""

                dias_solic  = nint(_XLS_COL_DIAS_SOL)
                dias_ccih   = nint(_XLS_COL_DIAS_CCIH)
                medicamento = nstr(_XLS_COL_MED)
                dias_em_uso = nint(_XLS_COL_DIAS_USO)
                medico      = nstr(_XLS_COL_MEDICO)

                if not medicamento:
                    continue

                pac = upsert_paciente(current_pront, current_nome, hospital)
                dt_inicio = _parse_date(dt_inicio_str)

                obj, created = ControleAtbRaw.objects.get_or_create(
                    paciente=pac,
                    medicamento=medicamento,
                    dt_inicio=dt_inicio,
                    defaults={
                        "acomodacao": current_acom,
                        "dias_solic":  dias_solic,
                        "dias_ccih":   dias_ccih,
                        "dias_em_uso": dias_em_uso,
                        "medico":      medico,
                    },
                )
                if created:
                    inserted += 1
                else:
                    changed = (
                        dias_em_uso > obj.dias_em_uso
                        or dias_solic > obj.dias_solic
                        or dias_ccih > obj.dias_ccih
                    )
                    if changed:
                        obj.dias_solic  = max(dias_solic,  obj.dias_solic)
                        obj.dias_ccih   = max(dias_ccih,   obj.dias_ccih)
                        obj.dias_em_uso = max(dias_em_uso, obj.dias_em_uso)
                        obj.save()
                        inserted += 1

    Importacao.objects.create(
        hospital=hospital, usuario=usuario,
        arquivo=path.name, tipo="controle_atb", registros=inserted,
    )
    return inserted


def importar_passagens(path: Path, hospital, usuario=None) -> int:
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    inserted = 0

    with transaction.atomic():
        for row_vals in rows[3:]:
            if not any(v for v in row_vals if v is not None):
                continue

            try:
                dias = int(float(row_vals[6])) if row_vals[6] is not None else 0
            except Exception:
                dias = 0
            if dias < 3:
                continue

            prontuario = _os_str(row_vals[4]) if row_vals[4] is not None else ""
            nome       = str(row_vals[3]).strip() if row_vals[3] is not None else ""
            dt_entrada_str = _excel_serial_to_date(row_vals[1]) if row_vals[1] is not None else ""
            dt_alta_str    = _excel_serial_to_date(row_vals[0]) if row_vals[0] is not None else ""

            if not prontuario or not dt_entrada_str:
                continue

            pac = upsert_paciente(prontuario, nome, hospital)
            dt_entrada = _parse_date(dt_entrada_str)
            dt_alta    = _parse_date(dt_alta_str)

            if not dt_entrada:
                continue

            obj, created = Internamento.objects.get_or_create(
                paciente=pac,
                dt_entrada=dt_entrada,
                defaults={"dt_alta": dt_alta, "dias": dias},
            )
            if created:
                inserted += 1
            else:
                changed = dias > obj.dias or dt_alta != obj.dt_alta
                if changed:
                    obj.dt_alta = dt_alta
                    obj.dias = max(dias, obj.dias)
                    obj.save()
                    inserted += 1

    Importacao.objects.create(
        hospital=hospital, usuario=usuario,
        arquivo=path.name, tipo="passagens", registros=inserted,
    )
    return inserted


def importar_arquivo(path, hospital, usuario=None) -> tuple:
    path = Path(path)
    tipo = _detect_type(path)
    if tipo == "microbiologia":
        n = importar_microbiologia(path, hospital, usuario)
        return tipo, n
    elif tipo == "controle_atb":
        n = importar_controle_atb(path, hospital, usuario)
        return tipo, n
    elif tipo == "passagens":
        n = importar_passagens(path, hospital, usuario)
        return tipo, n
    else:
        raise ValueError(f"Tipo de arquivo não reconhecido: {path.name}")


def importar_com_validacao(path, hospital, usuario, tipo_esperado: str) -> int:
    """Recusa o arquivo se o tipo detectado nao bater com tipo_esperado, em vez
    de importa-lo silenciosamente como outra coisa. Usado pelos comandos SSH
    dedicados (importar_microbiologia_ssh / importar_atb_ssh), que sao
    especificos por tipo e nao devem aceitar o arquivo errado."""
    path = Path(path)
    tipo = _detect_type(path)
    if tipo != tipo_esperado:
        raise ValueError(
            "Arquivo não reconhecido como '%s' (detectado: '%s')." % (tipo_esperado, tipo)
        )
    if tipo_esperado == "microbiologia":
        return importar_microbiologia(path, hospital, usuario)
    return importar_controle_atb(path, hospital, usuario)
